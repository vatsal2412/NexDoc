"""The naive baseline the structured adapter is measured against.

Flattens every sheet straight to CSV text: formulas are dropped in favor of
whatever openpyxl's data_only mode returns (which, for a file with no
cached values, is nothing — see formula.py's docstring), merged cells lose
their span, and cross-sheet relationships disappear entirely. This is
deliberately the "generic pipeline reduces everything to plain text" failure
mode described in CLAUDE.md's problem statement, not a strawman.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

from openpyxl import load_workbook


def flatten_to_csv(path: str | Path) -> str:
    """Returns one CSV block per sheet, concatenated with a sheet-name header
    line — the shape a human would get from "just export each tab to CSV."
    Formula cells come through as whatever cached value openpyxl can find,
    or blank if there is none.
    """
    wb = load_workbook(path, data_only=True)
    out = io.StringIO()
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        out.write(f"# sheet: {sheet_name}\n")
        writer = csv.writer(out)
        for row in ws.iter_rows():
            writer.writerow(["" if cell.value is None else cell.value for cell in row])
        out.write("\n")
    return out.getvalue()
