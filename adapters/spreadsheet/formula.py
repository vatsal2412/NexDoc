"""A deliberately small formula evaluator.

openpyxl never evaluates formulas, and a file that was authored by openpyxl
(rather than saved by Excel/LibreOffice) has no cached result to fall back
on either — `load_workbook(..., data_only=True)` returns None for every
formula cell in that case. So the adapter needs *some* way to produce a
resolved value and a real confidence signal.

Rather than pull in a full formula engine, this resolves exactly the
patterns Phase 1 needs to handle honestly:

    =SUM(A1:A9)                      same-sheet range
    ='Other Sheet'!A1                cross-sheet single-cell reference
    =SUM(A1:A9)+'Other Sheet'!A1     same-sheet range plus a cross-sheet cell

Anything else comes back as UNSUPPORTED rather than a guess — an adapter
that silently mis-evaluates an unfamiliar formula shape is worse than one
that flags it for review.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Union

from openpyxl.utils import range_boundaries
from openpyxl.workbook.workbook import Workbook

SUM_RANGE_RE = re.compile(r"^SUM\(([A-Z]+\d+):([A-Z]+\d+)\)$")
CROSS_SHEET_CELL_RE = re.compile(r"^(?:'([^']+)'|([A-Za-z0-9_ ]+))!([A-Z]+\d+)$")

UNRESOLVED_SHEET_NOT_FOUND = "sheet not found"
UNRESOLVED_UNSUPPORTED_PATTERN = "unsupported formula pattern"
UNRESOLVED_NON_NUMERIC_CELL = "range contains a non-numeric cell"


@dataclass
class Resolved:
    """Result of evaluating one formula cell."""

    value: Union[float, int, None]
    ok: bool
    reason: str | None  # set when ok is False
    same_sheet_refs: list[str]  # cell coordinates referenced within the formula's own sheet
    cross_sheet_refs: list[tuple[str, str]]  # (sheet_name, cell_coordinate) pairs


def _literal_value(ws, coordinate: str) -> Union[float, int, str, None]:
    return ws[coordinate].value


def _is_formula(value) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _sum_range(ws, start: str, end: str) -> tuple[Union[float, int, None], bool, list[str]]:
    min_col, min_row, max_col, max_row = range_boundaries(f"{start}:{end}")
    total = 0
    refs = []
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row=row, column=col)
            refs.append(cell.coordinate)
            value = cell.value
            if _is_formula(value):
                nested = evaluate(ws.parent, ws.title, cell.coordinate)
                if not nested.ok:
                    return None, False, refs
                value = nested.value
            if value is None:
                continue
            if not isinstance(value, (int, float)):
                return None, False, refs
            total += value
    return total, True, refs


def evaluate(workbook: Workbook, sheet_name: str, coordinate: str) -> Resolved:
    """Evaluate the formula at `sheet_name!coordinate`. Resolves recursively
    when a referenced cell is itself a formula (e.g. a cross-sheet total that
    is itself a SUM).
    """
    ws = workbook[sheet_name]
    raw = ws[coordinate].value
    if not _is_formula(raw):
        val = raw if isinstance(raw, (int, float)) else None
        return Resolved(val, val is not None, None, [], [])

    body = raw[1:].strip()

    m = SUM_RANGE_RE.match(body)
    if m:
        start, end = m.groups()
        total, ok, refs = _sum_range(ws, start, end)
        if not ok:
            return Resolved(None, False, UNRESOLVED_NON_NUMERIC_CELL, refs, [])
        return Resolved(total, True, None, refs, [])

    m = CROSS_SHEET_CELL_RE.match(body)
    if m:
        quoted_sheet, bare_sheet, cell = m.groups()
        target_sheet = quoted_sheet or bare_sheet
        if target_sheet not in workbook.sheetnames:
            return Resolved(None, False, UNRESOLVED_SHEET_NOT_FOUND, [], [(target_sheet, cell)])
        nested = evaluate(workbook, target_sheet, cell)
        if not nested.ok:
            return Resolved(None, False, nested.reason, [], [(target_sheet, cell)])
        return Resolved(nested.value, True, None, [], [(target_sheet, cell)])

    if "+" in body:
        left, right = body.rsplit("+", 1)
        left, right = left.strip(), right.strip()
        left_m = SUM_RANGE_RE.match(left)
        right_m = CROSS_SHEET_CELL_RE.match(right)
        if left_m and right_m:
            start, end = left_m.groups()
            left_total, left_ok, same_refs = _sum_range(ws, start, end)
            quoted_sheet, bare_sheet, cell = right_m.groups()
            target_sheet = quoted_sheet or bare_sheet
            if not left_ok:
                return Resolved(None, False, UNRESOLVED_NON_NUMERIC_CELL, same_refs, [(target_sheet, cell)])
            if target_sheet not in workbook.sheetnames:
                return Resolved(None, False, UNRESOLVED_SHEET_NOT_FOUND, same_refs, [(target_sheet, cell)])
            nested = evaluate(workbook, target_sheet, cell)
            if not nested.ok:
                return Resolved(None, False, nested.reason, same_refs, [(target_sheet, cell)])
            return Resolved(left_total + nested.value, True, None, same_refs, [(target_sheet, cell)])

    return Resolved(None, False, UNRESOLVED_UNSUPPORTED_PATTERN, [], [])
