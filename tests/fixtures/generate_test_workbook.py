"""Generates the synthetic .xlsx fixture used to exercise the spreadsheet adapter.

Produces tests/fixtures/budget_test.xlsx with two sheets:
  - "Q1 Plan": a title (merged cell), a header row, three data rows with
    in-sheet SUM formulas, and a totals row whose grand-total formula
    reaches across sheets into "Actuals".
  - "Actuals": same shape, referenced by "Q1 Plan"'s grand total.

This gives the adapter both confidence tiers (in-sheet formula vs.
cross-sheet formula) and both check types (cross-sheet reference resolution,
totals-vs-referenced-cells consistency) described in CLAUDE.md, in one small
file.
"""

from pathlib import Path

from openpyxl import Workbook

OUTPUT_PATH = Path(__file__).parent / "budget_test.xlsx"


def _write_budget_sheet(ws, title, data_rows):
    ws["A1"] = title
    ws.merge_cells("A1:E1")

    headers = ["Category", "Jan", "Feb", "Mar", "Total"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=2, column=col, value=header)

    first_data_row = 3
    for offset, (category, jan, feb, mar) in enumerate(data_rows):
        row = first_data_row + offset
        ws.cell(row=row, column=1, value=category)
        ws.cell(row=row, column=2, value=jan)
        ws.cell(row=row, column=3, value=feb)
        ws.cell(row=row, column=4, value=mar)
        ws.cell(row=row, column=5, value=f"=SUM(B{row}:D{row})")

    total_row = first_data_row + len(data_rows)
    last_data_row = total_row - 1
    ws.cell(row=total_row, column=1, value="Total")
    ws.cell(row=total_row, column=2, value=f"=SUM(B{first_data_row}:B{last_data_row})")
    ws.cell(row=total_row, column=3, value=f"=SUM(C{first_data_row}:C{last_data_row})")
    ws.cell(row=total_row, column=4, value=f"=SUM(D{first_data_row}:D{last_data_row})")
    return total_row, first_data_row, last_data_row


def build_workbook():
    wb = Workbook()

    ws_actuals = wb.active
    ws_actuals.title = "Actuals"
    _write_budget_sheet(
        ws_actuals,
        "Q1 Actuals",
        [
            ("Marketing", 1150, 1380, 1250),
            ("Engineering", 2950, 3050, 2900),
            ("Sales", 880, 920, 990),
        ],
    )
    # Actuals!E6 = SUM(E3:E5), a plain in-sheet total — no cross-sheet part.
    ws_actuals.cell(row=6, column=5, value="=SUM(E3:E5)")

    ws_plan = wb.create_sheet("Q1 Plan")
    _write_budget_sheet(
        ws_plan,
        "Q1 Budget Plan",
        [
            ("Marketing", 1200, 1400, 1300),
            ("Engineering", 3000, 3100, 2950),
            ("Sales", 900, 950, 1000),
        ],
    )
    # Q1 Plan!E6: grand total = sum of this sheet's own row totals PLUS the
    # Actuals grand total — the cross-sheet formula the adapter must flag.
    ws_plan.cell(row=6, column=5, value="=SUM(E3:E5)+'Actuals'!E6")

    # Deliberately broken: a link to a sheet that was renamed/deleted, so the
    # cross-sheet-reference check has a real failing instance to catch, not
    # just a passing one.
    ws_plan.cell(row=2, column=6, value="Forecast link")
    ws_plan.cell(row=3, column=6, value="='Forecast Q4'!B3")

    wb.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_workbook()
    print(f"Wrote {path}")
